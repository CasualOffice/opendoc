# SPDX-License-Identifier: Apache-2.0
"""One Python measurement in a fresh process, for benchmarks/sdk/run.mjs.

    python python_task.py <task> <library> <input> <iterations>

The same python-docx program runs against python-docx and against opendoc,
with only the import changed, which is the claim being measured. Prints one
JSON line: median milliseconds and peak RSS in MB.
"""
import json
import resource
import sys
import time

task, library, value, iterations = sys.argv[1], sys.argv[2], sys.argv[3], int(sys.argv[4])

if library == "python-docx":
    from docx import Document
else:
    from opendoc import Document


def generate(n: int) -> int:
    document = Document()
    document.add_heading("Benchmark report", 1)
    for i in range(n):
        p = document.add_paragraph("Paragraph %d: the quick brown fox jumps over the lazy dog. " % i)
        p.add_run("Bold").bold = True
    table = document.add_table(rows=1, cols=3)
    for r in range(max(10, n // 10)):
        cells = table.add_row().cells
        cells[0].text = "r%dc0" % r
    import io

    out = io.BytesIO()
    document.save(out)
    return len(out.getvalue())


def read(path: str) -> int:
    document = Document(path)
    return sum(len(p.text) for p in document.paragraphs)


def run() -> int:
    if task == "generate":
        return generate(int(value))
    return read(value)


output = run()
times = []
for _ in range(iterations):
    start = time.perf_counter()
    output = run()
    times.append((time.perf_counter() - start) * 1000)
times.sort()
peak = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
peak_mb = peak / (1024 * 1024) if sys.platform == "darwin" else peak / 1024
print(json.dumps({"ms": round(times[len(times) // 2], 2), "peakMB": round(peak_mb, 1), "output": output}))
